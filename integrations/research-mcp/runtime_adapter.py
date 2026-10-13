"""Private Lingbi adapter: preserve provenance when adapting PDF/text to PaperCheck.

Upstream's evidence extractor only accepts DOCX. Conversion runs in the isolated
MCP service account, never in the web server or against another user's files.
"""
import hashlib
import os
from pathlib import Path
import sys

import uvicorn
from sci_mcp import factory
from sci_mcp.workspace import workspace_dir, resolve_under

original_add = factory._add_tool


def add_tool(server, name, description, function, enums=None):
    if server.spec.id == "papercheck" and name == "extract_citation_evidence":
        original_extract = function

        def extract(workspace_file: str) -> dict:
            root = workspace_dir("papercheck")
            source = resolve_under(root, workspace_file)
            if not source.is_file() or source.stat().st_size > 6 * 1024 * 1024:
                raise ValueError("输入文件不存在或超过 6MB。")
            extension = source.suffix.lower()
            if extension == ".docx":
                return original_extract(workspace_file)
            if extension not in {".pdf", ".txt", ".md"}:
                raise ValueError("仅支持 PDF、DOCX 或 UTF-8 文本文件。")
            from docx import Document
            document = Document()
            document.add_heading("Citation extraction source", 0)
            characters = 0
            lines = 0

            def append(text):
                nonlocal characters, lines
                characters += len(text)
                if characters > 1_000_000:
                    raise ValueError("文本过长，请拆分论文后再核验。")
                for line in text.splitlines():
                    line = line.strip()
                    if not line:
                        continue
                    lines += 1
                    if lines > 20_000:
                        raise ValueError("文本行数超过核验限制。")
                    if line.casefold() in {"references", "bibliography", "参考文献"}:
                        document.add_heading(line, 1)
                    else:
                        document.add_paragraph(line)

            if extension == ".pdf":
                import fitz
                with fitz.open(source) as pdf:
                    if pdf.is_encrypted or pdf.page_count > 200:
                        raise ValueError("请上传未加密、最多 200 页的 PDF。")
                    for index, page in enumerate(pdf):
                        document.add_paragraph(f"【PDF 第 {index + 1} 页】")
                        append(page.get_text("text"))
                if characters < 20:
                    raise ValueError("该 PDF 缺少可提取文本，扫描件需先完成 OCR。")
                source_format = "pdf-text-layer"
            else:
                append(source.read_text(encoding="utf-8-sig"))
                source_format = "utf-8-text"
            converted = source.with_name(source.stem + ".converted.docx")
            document.save(converted)
            result = original_extract(converted.relative_to(root).as_posix())
            return {**result, "source_format": source_format,
                    "original_file": source.relative_to(root).as_posix(),
                    "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                    "note": "仅做文本层证据抽取；PDF 页码请回原文件核对，不代表正文主张已获支持。"}

        function = extract
        description = "抽取 DOCX、PDF 文本层或 UTF-8 文本中的引文与参考文献；扫描 PDF 需先 OCR。"
    return original_add(server, name, description, function, enums)


factory._add_tool = add_tool

if __name__ == "__main__":
    from sci_mcp.asgi import create_product_app
    uvicorn.run(create_product_app(sys.argv[1]), host="127.0.0.1", port=int(os.environ["PORT"]), access_log=False, log_level="warning")
